"""Download an assessment report as PDF or Word.

Both formats are built from one plain ``ReportData`` structure, so the
content is identical. PDF uses PyMuPDF's Story API (HTML -> PDF) and Word
uses python-docx; both libraries are already project dependencies.
"""

from __future__ import annotations

import html
import io
import re
from dataclasses import dataclass, field
from datetime import datetime


SEVERITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3, "unrated": 4}

COVERAGE_LABELS = {
    "assessed": "Assessed",
    "insufficient_evidence": "Insufficient evidence",
    "no_evidence": "No evidence",
    "not_reported": "Not reported",
}


def label(value: str | None) -> str:
    text = (value or "").replace("_", " ").strip().title()
    return re.sub(r"\bAi\b", "AI", text)


@dataclass
class ReportFinding:
    title: str
    category: str
    severity: str
    finding_type: str
    description: str
    rule_id: str | None = None
    evidence_count: int = 0


@dataclass
class ReportCoverage:
    category: str
    status: str
    finding_count: int = 0
    evidence_chunks: int = 0
    note: str = ""


@dataclass
class ReportData:
    vendor_name: str
    assessment_id: str
    completed_at: datetime | None
    overall_risk: str | None
    overall_score: float | None
    coverage_percent: float | None
    policy_version: str | None
    summary: str
    findings: list[ReportFinding] = field(default_factory=list)
    coverage: list[ReportCoverage] = field(default_factory=list)

    def sorted_findings(self) -> list[ReportFinding]:
        return sorted(
            self.findings,
            key=lambda f: (SEVERITY_ORDER.get(f.severity, 9), f.category, f.title),
        )

    def headline(self) -> list[tuple[str, str]]:
        return [
            ("Overall risk", label(self.overall_risk or "unrated")),
            (
                "Risk score",
                f"{round(self.overall_score)} / 100"
                if self.overall_score is not None
                else "Not calculated",
            ),
            (
                "Coverage",
                f"{round(self.coverage_percent)}%"
                if self.coverage_percent is not None
                else "Not measured",
            ),
            ("Findings", str(len(self.findings))),
            (
                "Completed",
                self.completed_at.strftime("%B %d, %Y %H:%M UTC")
                if self.completed_at
                else "Unknown",
            ),
            ("Policy", self.policy_version or "Not applied"),
        ]


def report_from_response(response, *, vendor_name: str) -> ReportData:
    """Build ReportData from an AssessmentResponse."""
    metadata = response.metadata or {}
    findings = [
        ReportFinding(
            title=item.title,
            category=item.category,
            severity=item.severity or "unrated",
            finding_type=(item.metadata or {}).get("finding_type", "other"),
            description=item.description,
            rule_id=item.rule_id,
            evidence_count=len((item.metadata or {}).get("evidence") or []),
        )
        for item in response.findings
    ]
    coverage = [
        ReportCoverage(
            category=entry.get("category", ""),
            status=entry.get("status", "not_reported"),
            finding_count=int(entry.get("finding_count") or 0),
            evidence_chunks=int(entry.get("evidence_chunks") or 0),
            note=entry.get("note") or "",
        )
        for entry in (metadata.get("category_coverage") or [])
        if isinstance(entry, dict)
    ]
    return ReportData(
        vendor_name=vendor_name,
        assessment_id=str(response.id),
        completed_at=response.completed_at,
        overall_risk=response.overall_risk,
        overall_score=response.overall_score,
        coverage_percent=metadata.get("coverage_percent"),
        policy_version=metadata.get("risk_policy_version"),
        summary=metadata.get("analysis_summary") or "",
        findings=findings,
        coverage=coverage,
    )


def export_filename(report: ReportData, extension: str) -> str:
    vendor = re.sub(r"[^A-Za-z0-9]+", "-", report.vendor_name).strip("-").lower()
    date = report.completed_at.strftime("%Y-%m-%d") if report.completed_at else "report"
    return f"{vendor or 'vendor'}-risk-assessment-{date}.{extension}"


DISCLAIMER = (
    "Findings are generated from the vendor's own documents by an AI model and "
    "rated by deterministic policy rules. A category without findings is not "
    "proof of a strong control; review categories marked no or insufficient "
    "evidence before relying on this report."
)


# ------------------------------------------------------------
# Word
# ------------------------------------------------------------

def build_docx(report: ReportData) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)

    document.add_heading(f"Vendor Risk Assessment: {report.vendor_name}", level=0)

    table = document.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for name, value in report.headline():
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = value

    if report.summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(report.summary)

    document.add_heading("Findings", level=1)
    if not report.findings:
        document.add_paragraph("No findings were generated.")
    for finding in report.sorted_findings():
        document.add_heading(finding.title, level=2)
        meta = document.add_paragraph()
        run = meta.add_run(
            f"{label(finding.severity)} · {label(finding.category)} · "
            f"{label(finding.finding_type)}"
            + (f" · Rule {finding.rule_id}" if finding.rule_id else "")
            + f" · {finding.evidence_count} evidence source(s)"
        )
        run.bold = True
        run.font.color.rgb = RGBColor(0x47, 0x54, 0x67)
        document.add_paragraph(finding.description)

    if report.coverage:
        document.add_heading("Assessment coverage", level=1)
        coverage = document.add_table(rows=1, cols=4)
        coverage.style = "Light List Accent 1"
        header = coverage.rows[0].cells
        for cell, text in zip(header, ["Category", "Coverage", "Findings", "Note"]):
            cell.text = text
        for entry in report.coverage:
            cells = coverage.add_row().cells
            cells[0].text = label(entry.category)
            cells[1].text = COVERAGE_LABELS.get(entry.status, label(entry.status))
            cells[2].text = str(entry.finding_count)
            cells[3].text = entry.note

    note = document.add_paragraph()
    note_run = note.add_run(DISCLAIMER)
    note_run.italic = True
    note_run.font.size = Pt(8.5)
    document.add_paragraph(f"Assessment ID: {report.assessment_id}").runs[0].font.size = Pt(8.5)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------
# PDF
# ------------------------------------------------------------

PDF_CSS = """
* { font-family: sans-serif; }
body { font-size: 10pt; color: #101828; }
h1 { font-size: 18pt; margin-bottom: 4pt; }
h2 { font-size: 13pt; margin-top: 14pt; border-bottom: 1px solid #e4e7ec; }
h3 { font-size: 11pt; margin-top: 10pt; margin-bottom: 2pt; }
p { margin-top: 2pt; margin-bottom: 4pt; }
.meta { color: #475467; font-size: 9pt; font-weight: bold; }
.headline { color: #344054; }
.small { color: #667085; font-size: 8pt; font-style: italic; }
"""


def build_report_html(report: ReportData) -> str:
    e = html.escape
    parts = [f"<h1>Vendor Risk Assessment: {e(report.vendor_name)}</h1>"]
    parts.append(
        "<p class='headline'>"
        + "<br/>".join(f"<b>{e(name)}:</b> {e(value)}" for name, value in report.headline())
        + "</p>"
    )
    if report.summary:
        parts.append(f"<h2>Summary</h2><p>{e(report.summary)}</p>")

    parts.append("<h2>Findings</h2>")
    if not report.findings:
        parts.append("<p>No findings were generated.</p>")
    for finding in report.sorted_findings():
        meta = (
            f"{label(finding.severity)} · {label(finding.category)} · "
            f"{label(finding.finding_type)}"
            + (f" · Rule {finding.rule_id}" if finding.rule_id else "")
            + f" · {finding.evidence_count} evidence source(s)"
        )
        parts.append(
            f"<h3>{e(finding.title)}</h3>"
            f"<p class='meta'>{e(meta)}</p>"
            f"<p>{e(finding.description)}</p>"
        )

    if report.coverage:
        parts.append("<h2>Assessment coverage</h2>")
        for entry in report.coverage:
            status = COVERAGE_LABELS.get(entry.status, label(entry.status))
            parts.append(
                f"<p><b>{e(label(entry.category))}</b> — {e(status)}, "
                f"{entry.finding_count} finding(s)"
                + (f"<br/>{e(entry.note)}" if entry.note else "")
                + "</p>"
            )

    parts.append(f"<p class='small'>{e(DISCLAIMER)}</p>")
    parts.append(f"<p class='small'>Assessment ID: {e(report.assessment_id)}</p>")
    return "<html><body>" + "".join(parts) + "</body></html>"


def build_pdf(report: ReportData) -> bytes:
    import pymupdf

    buffer = io.BytesIO()
    story = pymupdf.Story(html=build_report_html(report), user_css=PDF_CSS)
    writer = pymupdf.DocumentWriter(buffer)
    mediabox = pymupdf.paper_rect("letter")
    content_area = mediabox + (54, 54, -54, -54)

    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(content_area)
        story.draw(device)
        writer.end_page()

    writer.close()
    return buffer.getvalue()
