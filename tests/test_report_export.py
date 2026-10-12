"""Report content shared by the PDF and Word exports."""

from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest

from vendor_risk_analyzer.assessments.export import (
    ReportCoverage,
    ReportData,
    ReportFinding,
    build_docx,
    build_report_html,
    export_filename,
)


pytestmark = pytest.mark.unit


def report(**changes):
    values = dict(
        vendor_name="Larkspur Data Systems",
        assessment_id="assessment-1",
        completed_at=datetime(2026, 10, 12, 4, 0, tzinfo=timezone.utc),
        overall_risk="high",
        overall_score=75.7,
        coverage_percent=88.9,
        policy_version="vendor-risk-policy-v3",
        summary="Summary with <script>alert(1)</script>",
        findings=[
            ReportFinding("Breach notice conflict", "incident_response", "medium",
                          "contradiction", "72h vs 24h", "VRP-IR-CONTRADICTION-001", 2),
            ReportFinding("MFA not enforced", "access_control", "high",
                          "explicit_risk", "14 contractor accounts", "VRP-AC-EXPLICIT-001", 1),
        ],
        coverage=[ReportCoverage("ai_governance", "insufficient_evidence", 0, 6, "Not mentioned")],
    )
    values.update(changes)
    return ReportData(**values)


def test_headline_and_ordering():
    data = report()

    assert dict(data.headline())["Risk score"] == "76 / 100"
    assert dict(data.headline())["Coverage"] == "89%"
    assert [f.title for f in data.sorted_findings()] == [
        "MFA not enforced",
        "Breach notice conflict",
    ]
    assert dict(report(overall_score=None, coverage_percent=None).headline())[
        "Risk score"
    ] == "Not calculated"


def test_pdf_html_escapes_model_text():
    html = build_report_html(report())

    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "AI Governance" in html


def test_word_report_contains_findings_and_coverage():
    from docx import Document

    document = Document(io.BytesIO(build_docx(report())))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]

    assert "Vendor Risk Assessment: Larkspur Data Systems" in text
    assert text.index("MFA not enforced") < text.index("Breach notice conflict")
    assert "76 / 100" in cells
    assert "Insufficient evidence" in cells


def test_filename_is_safe():
    assert export_filename(report(vendor_name='A/B "Vendor" Inc.'), "pdf") == (
        "a-b-vendor-inc-risk-assessment-2026-10-12.pdf"
    )
    assert export_filename(report(completed_at=None), "docx").endswith("-report.docx")
